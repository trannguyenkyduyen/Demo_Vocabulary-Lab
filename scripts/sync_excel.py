import io, os, json, re, sys, zipfile
from urllib.request import Request, urlopen
from openpyxl import load_workbook

SHARE_URL = os.environ.get('ONEDRIVE_SHARE_URL', '').strip()
if not SHARE_URL:
    raise SystemExit('Missing ONEDRIVE_SHARE_URL GitHub Actions secret')
# OneDrive Personal public share links commonly support a download=1 query flag.
sep = '&' if '?' in SHARE_URL else '?'
url = SHARE_URL + sep + 'download=1'
req = Request(url, headers={'User-Agent': 'Mozilla/5.0 VocabularyLabSync/1.0'})
with urlopen(req, timeout=45) as response:
    raw = response.read(15_000_000)
if not zipfile.is_zipfile(io.BytesIO(raw)):
    raise SystemExit('OneDrive returned a webpage rather than an XLSX file. Check public sharing or download link. No changes committed.')
wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
if 'Vocabulary' not in wb.sheetnames:
    raise SystemExit('Missing Vocabulary sheet')
ws = wb['Vocabulary']
headers = [str(c.value or '').strip() for c in next(ws.iter_rows(min_row=1,max_row=1))]
expected = ['ID','English','Nghĩa tiếng Việt','Nhóm từ vựng','Từ loại','Example (EN)','Ví dụ (VI)']
if headers[:7] != expected:
    raise SystemExit('Unexpected Excel columns: '+repr(headers))
words = []
for cells in ws.iter_rows(min_row=2,values_only=True):
    if not any(x is not None and str(x).strip() for x in cells[:7]): continue
    v = [(str(x).strip() if x is not None else '') for x in cells[:7]]
    if not all(v[i] for i in [0,1,2,3]):
        raise SystemExit('Missing required field in Excel row '+str(len(words)+2))
    words.append(dict(zip(['id','en','vi','group','pos','example','exampleVi'],v)))
if len(words)<10 or len(set(w['id'] for w in words)) != len(words):
    raise SystemExit('Too few words or duplicate IDs; refusing to publish')

journey_extra={}
if 'My Life Journey' in wb.sheetnames:
    for cells in wb['My Life Journey'].iter_rows(min_row=2,values_only=True):
        if len(cells)<8 or not cells[3]: continue
        en=str(cells[3]).strip()
        journey_extra[en]={'en':en,'vi':str(cells[5] or ''),'example':str(cells[6] or ''),'exampleVi':str(cells[7] or '')}

payload={'id':'grade12-life-stories-v1','title':'Topic 01 – Life Stories','words':words}
# JS rather than JSON enables same-origin static hosting without CORS issues.
js='window.TOPIC01_DATA='+json.dumps(payload,ensure_ascii=False,separators=(',',':'))+';\n'
js+='window.TOPIC01_JOURNEY='+json.dumps(journey_extra,ensure_ascii=False,separators=(',',':'))+';\n'
with open('topic01-data.js','w',encoding='utf-8') as f:f.write(js)
print(f'Updated {len(words)} vocabulary entries and {len(journey_extra)} journey entries')
